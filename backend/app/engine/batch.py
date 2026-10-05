"""Чтение файлов CSV/XLSX, проверка формата данных и расчёт по каждой строке.

Проверка формата идёт в два уровня:
1. Весь файл (фатальные ошибки, файл не обрабатывается): формат и кодировка, пустой файл, нет обязательных
   столбцов, столбцы с персональными данными, повторяющиеся показатели, нераспознанные единицы в заголовках.
2. Каждая ячейка (ошибка строки, остальные строки обрабатываются): не число, значение вне правдоподобного
   диапазона (часто это перепутанные единицы), нераспознанный пол, неверный возраст, некорректный ID.
Пустой пол или гемоглобин в строке — не ошибка формата: такую строку ML-модуль вернёт со статусом
insufficient_data (достаточность данных проверяет только модуль).
Каждая проблема описывается словарём {level, row, column, message}, чтобы интерфейс показал, где именно ошибка.
"""
import csv
import io
import re
from dataclasses import dataclass, field
from typing import Any

from . import analytes as an
from . import labels as lb
from .analytes import normalize_key
from .pipeline import screen
from .predictor import Predictor, model_info
from .validation import CaseInput, validate_case

_ID_COLUMNS = {"case_id", "patient_id", "id", "id_случая"}
_AGE_COLUMNS = {"age", "age_years", "возраст"}
_SEX_COLUMNS = {"sex", "gender", "пол"}
# Готовые ответы (целевые переменные). Их наличие допустимо, но в расчёте они не используются.
_LABEL_COLUMNS = {
    "anemia", "iron_deficiency", "b12_deficiency", "folate_deficiency", "b6_deficiency", "copper_deficiency",
    "inflammation_anemia", "mixed_deficiency", "anemia_class", "deficiency_cause",
}
# Столбцы с персональными данными: такой файл не принимается (принцип минимизации данных)
_PERSONAL_COLUMNS = {
    "fio", "full_name", "fullname", "name", "first_name", "last_name", "surname", "patronymic", "middle_name",
    "phone", "telephone", "email", "e_mail", "snils", "passport", "address", "inn", "oms", "polis", "birth_date",
    "фио", "фамилия", "имя", "отчество", "телефон", "адрес", "снилс", "паспорт", "инн", "полис", "дата_рождения",
}
# «hemoglobin (g/dL)», «ferritin [ng/mL]», «Гемоглобин, г/дл»
_HEADER_UNIT_RE = re.compile(r"^\s*(.+?)\s*(?:[\(\[]\s*(.+?)\s*[\)\]]|,\s*(.+?))\s*$")
_MAX_ISSUES = 300


def plural(n: int, one: str, few: str, many: str) -> str:
    """1 строка, 2 строки, 5 строк."""
    if n % 10 == 1 and n % 100 != 11:
        return one
    if 2 <= n % 10 <= 4 and not 12 <= n % 100 <= 14:
        return few
    return many


def issue(level: str, message: str, row: int | None = None, column: str | None = None) -> dict:
    return {"level": level, "row": row, "column": column, "message": message}


class BatchFileError(Exception):
    """Файл не может быть обработан целиком. `issues` — все найденные проблемы уровня файла."""

    def __init__(self, message: str, issues: list[dict] | None = None):
        super().__init__(message)
        self.issues = issues or [issue("error", message)]


def _decode(content: bytes) -> str:
    for encoding in ("utf-8-sig", "cp1251"):
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise BatchFileError("Не удалось прочитать кодировку файла. Сохраните CSV в кодировке UTF-8.")


def _read_csv(content: bytes) -> list[list[Any]]:
    text = _decode(content)
    first_line = text.splitlines()[0] if text.strip() else ""
    delimiter = max((",", ";", "\t"), key=first_line.count)  # разделитель определяем по строке заголовков
    return [list(row) for row in csv.reader(io.StringIO(text), delimiter=delimiter)]


def _read_xlsx(content: bytes) -> list[list[Any]]:
    from openpyxl import load_workbook

    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001 — любой сбой чтения превращаем в понятную ошибку
        raise BatchFileError("Не удалось прочитать XLSX-файл. Проверьте, что файл не повреждён.") from exc
    sheet = workbook.worksheets[0]
    return [["" if c is None else c for c in row] for row in sheet.iter_rows(values_only=True)]


def read_table(filename: str, content: bytes) -> list[list[Any]]:
    name = (filename or "").lower()
    if not content:
        raise BatchFileError("Файл пустой.")
    if name.endswith(".csv"):
        table = _read_csv(content)
    elif name.endswith(".xlsx"):
        table = _read_xlsx(content)
    else:
        raise BatchFileError("Поддерживаются только файлы CSV и XLSX.")
    table = [row for row in table if any(str(c).strip() for c in row)]
    if len(table) < 2:
        raise BatchFileError("В файле нет данных: нужны строка заголовков и хотя бы одна строка пациента.")
    return table


def _split_header(raw: Any) -> tuple[str, str | None]:
    """Отделяет единицу измерения от названия столбца, если она указана в заголовке."""
    text = str(raw)
    if an.resolve_key(text) or normalize_key(text) in _ID_COLUMNS | _AGE_COLUMNS | _SEX_COLUMNS:
        return text, None
    match = _HEADER_UNIT_RE.match(text)
    if match:
        return match.group(1), match.group(2) or match.group(3)
    return text, None


@dataclass
class TableLayout:
    """Что лежит в каждом столбце файла."""
    roles: list[tuple[str, str | None]] = field(default_factory=list)  # (роль, ключ показателя)
    column_names: dict[str, str] = field(default_factory=dict)         # ключ показателя/поля -> заголовок в файле
    units: dict[str, str] = field(default_factory=dict)
    ignored: list[str] = field(default_factory=list)
    labels: list[str] = field(default_factory=list)
    warnings: list[dict] = field(default_factory=list)


def inspect_header(header: list[Any]) -> TableLayout:
    """Проверяет строку заголовков. Фатальные проблемы собираются все сразу и выбрасываются одним исключением."""
    layout = TableLayout()
    errors: list[dict] = []
    seen: dict[str, str] = {}
    personal: list[str] = []
    for raw in header:
        raw_text = str(raw).strip()
        name, unit = _split_header(raw_text)
        key = normalize_key(name)
        if raw_text == "":
            layout.roles.append(("skip", None))
        elif key in _PERSONAL_COLUMNS:
            personal.append(raw_text)
            layout.roles.append(("skip", None))
        elif key in _ID_COLUMNS | _AGE_COLUMNS | _SEX_COLUMNS:
            role = "id" if key in _ID_COLUMNS else "age" if key in _AGE_COLUMNS else "sex"
            if role in seen:
                errors.append(issue("error", f"Столбец «{raw_text}» дублирует «{seen[role]}».", column=raw_text))
            seen[role] = raw_text
            layout.column_names[{"id": "case_id"}.get(role, role)] = raw_text
            layout.roles.append((role, None))
        elif key in _LABEL_COLUMNS:
            layout.labels.append(raw_text)
            layout.roles.append(("skip", None))
        elif (analyte_key := an.resolve_key(name)) is not None:
            if analyte_key in seen:
                errors.append(issue("error", f"Показатель {an.BY_KEY[analyte_key].label} указан в двух столбцах: "
                                             f"«{seen[analyte_key]}» и «{raw_text}». Оставьте один.", column=raw_text))
            seen[analyte_key] = raw_text
            layout.column_names[f"laboratory_data.{analyte_key}"] = raw_text
            layout.roles.append(("lab", analyte_key))
            if unit:
                if an.unit_factor(an.BY_KEY[analyte_key], unit) is None:
                    errors.append(issue("error", f"Единица «{unit}» не поддерживается для показателя "
                                                 f"{an.BY_KEY[analyte_key].label}. Основная единица: {an.BY_KEY[analyte_key].unit}.",
                                        column=raw_text))
                layout.units[analyte_key] = unit
        else:
            layout.ignored.append(raw_text)
            layout.roles.append(("skip", None))

    if personal:
        errors.append(issue("error", "Файл содержит столбцы с персональными данными (" + ", ".join(personal) + "). "
                                     "Удалите их: сервису нужен только псевдонимизированный ID."))
    present = {r for r, _ in layout.roles}
    for needed, title in (("id", "patient_id (или case_id)"), ("age", "age_years (или age)"), ("sex", "sex")):
        if needed not in present:
            errors.append(issue("error", f"Нет обязательного столбца {title}."))
    if "hemoglobin" not in {k for _, k in layout.roles}:
        errors.append(issue("error", "Нет обязательного столбца hemoglobin (гемоглобин)."))
    if errors:
        raise BatchFileError(errors[0]["message"], errors)

    if layout.ignored:
        layout.warnings.append(issue("warning", "Столбцы не распознаны и не учтены: " + ", ".join(layout.ignored) + "."))
    if layout.labels:
        layout.warnings.append(issue("warning", "Столбцы с готовыми ответами не используются в расчёте: "
                                                + ", ".join(layout.labels) + "."))
    for key, unit in layout.units.items():
        layout.warnings.append(issue("warning", f"Значения будут переведены из «{unit}» в {an.BY_KEY[key].unit}.",
                                     column=layout.column_names[f"laboratory_data.{key}"]))
    return layout


def _row_values(layout: TableLayout, raw_row: list[Any]) -> tuple[dict[str, Any], dict[str, Any]]:
    values: dict[str, Any] = {}
    labs: dict[str, Any] = {}
    for index, (role, key) in enumerate(layout.roles):
        cell = raw_row[index] if index < len(raw_row) else ""
        if role in ("id", "age", "sex"):
            values[role] = cell
        elif role == "lab" and key is not None and str(cell).strip() != "":
            labs[key] = cell
    return values, labs


def check_rows(layout: TableLayout, data_rows: list[list[Any]]):
    """Проверяет каждую строку. Возвращает список (номер, ID, CaseInput|None, ошибки, предупреждения, переводы единиц)."""
    checked = []
    seen_ids: dict[str, int] = {}
    for number, raw_row in enumerate(data_rows, start=1):
        values, labs = _row_values(layout, raw_row)
        check = validate_case(values.get("id"), values.get("age"), values.get("sex"), labs, layout.units)
        row_errors = [issue("error", e.message, number, layout.column_names.get(e.field, e.field)) for e in check.errors]
        row_warnings = []
        case_id = str(values.get("id", "")).strip() or None
        if case_id:
            if case_id in seen_ids:
                row_warnings.append(issue("warning", f"ID «{case_id}» уже встречался в строке {seen_ids[case_id]}.",
                                          number, layout.column_names.get("case_id")))
            else:
                seen_ids[case_id] = number
        checked.append((number, case_id, check.case, row_errors, row_warnings, check.conversions))
    return checked


def _file_check(layout: TableLayout, rows_total: int, issues_list: list[dict]) -> dict:
    ordered = sorted(issues_list, key=lambda i: (i["level"] != "error", i["row"] or 0))
    errors = sum(1 for i in ordered if i["level"] == "error")
    return {
        "ok": errors == 0,
        "rows": rows_total,
        "recognized_columns": sum(1 for r, _ in layout.roles if r == "lab"),
        "errors": errors,
        "warnings": len(ordered) - errors,
        "issues": ordered[:_MAX_ISSUES],
        "issues_truncated": max(0, len(ordered) - _MAX_ISSUES),
    }


def process_batch(table: list[list[Any]], predictor: Predictor, max_rows: int) -> dict:
    header, data_rows = table[0], table[1:]
    if len(data_rows) > max_rows:
        raise BatchFileError(f"В файле {len(data_rows)} {plural(len(data_rows), 'строка', 'строки', 'строк')}, "
                             f"максимум для одной загрузки: {max_rows}.")
    layout = inspect_header(header)
    all_issues = list(layout.warnings)

    checked = check_rows(layout, data_rows)
    valid_cases = [item[2] for item in checked if item[2] is not None]
    ml_results = iter(predictor.predict_many(valid_cases))   # один вызов модели на весь файл

    rows: list[dict] = []
    for number, case_id, case, row_errors, row_warnings, _conv in checked:
        all_issues += row_errors + row_warnings
        row: dict = {"row": number, "case_id": case_id, "status": "error",
                     "errors": [e["message"] for e in row_errors], "warnings": [w["message"] for w in row_warnings]}
        if case is not None:
            result = screen(case, predictor, next(ml_results))
            ml = result["ml_result"]
            row.update(
                status="completed", sex=case.sex, age=case.age,
                screening_status=ml["status"], anemia=result["anemia"]["detected"],
                hemoglobin=result["anemia"]["hemoglobin"],
                deficiency_cause=ml["deficiency_cause"], deficiency_cause_label=result["deficiency_cause_label"],
                anemia_class=ml["anemia_class"], anemia_class_label=result["anemia_class_label"],
                model_score=ml["model_score"], title=result["report"]["title"], severity=result["report"]["severity"],
                # нормализованный вход (стандартные единицы): по нему открывается полный результат случая
                input={"case_id": case.case_id, "age": case.age, "sex": case.sex, "laboratory_data": case.labs},
            )
        rows.append(row)

    completed = [r for r in rows if r["status"] == "completed"]
    return {
        "summary": {
            "total": len(rows), "completed": len(completed), "errors": len(rows) - len(completed),
            "anemia_detected": sum(1 for r in completed if r["anemia"]),
            "with_deficiency": sum(1 for r in completed if r["deficiency_cause"] in lb.NUTRITIONAL_CAUSES),
            "insufficient_data": sum(1 for r in completed if r["screening_status"] == "insufficient_data"),
        },
        "rows": rows,
        "ignored_columns": layout.ignored + layout.labels,
        "units_from_header": layout.units,
        "file_check": _file_check(layout, len(rows), all_issues),
        "model": model_info(predictor),
    }


def process_single(table: list[list[Any]], predictor: Predictor) -> tuple[dict | None, dict, CaseInput | None]:
    """Файл одного пациента: ровно одна строка данных. Возвращает (результат | None, file_check, случай)."""
    header, data_rows = table[0], table[1:]
    if len(data_rows) != 1:
        raise BatchFileError(
            f"В файле {len(data_rows)} {plural(len(data_rows), 'строка', 'строки', 'строк')} с пациентами, "
            "а для этого режима нужна ровно одна. "
            "Для нескольких пациентов откройте режим «Много пациентов».")
    layout = inspect_header(header)
    (_number, _cid, case, row_errors, row_warnings, conversions), = check_rows(layout, data_rows)
    check = _file_check(layout, 1, layout.warnings + row_errors + row_warnings)
    if case is None:
        return None, check, None
    result = screen(case, predictor)
    result["unit_conversions"] = conversions
    return result, check, case
